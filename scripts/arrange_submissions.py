#!/usr/bin/env python3
"""
Arrange WEEK2/WEEK3/WEEK4 submission folders to match WEEK1 naming,
generate architecture JPGs, and build Week 1-style progress reports plus
Week 4 assignment documents.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor  # noqa: F401

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SUBMISSION = PROJECT_ROOT / "submission"

# Import shared helpers from the existing builder
import sys

sys.path.insert(0, str(PROJECT_ROOT / "scripts"))
from build_submission_docx import (  # noqa: E402
    add_page_numbers,
    add_title_block,
    parse_markdown_to_doc,
)


# ---------------------------------------------------------------------------
# Fonts
# ---------------------------------------------------------------------------

def _font(size: int, bold: bool = False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


# ---------------------------------------------------------------------------
# Box-and-arrow architecture diagrams (Week 1 style)
# ---------------------------------------------------------------------------

def _draw_box(draw, xy, text, font, fill="#ffffff", outline="#222222", radius=6):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=2)
    cx = (x1 + x2) / 2
    cy = (y1 + y2) / 2
    bbox = draw.textbbox((0, 0), text, font=font)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.text((cx - tw / 2, cy - th / 2), text, fill="#111111", font=font)


def _arrow(draw, start, end):
    draw.line([start, end], fill="#333333", width=2)
    x, y = end
    if end[1] > start[1]:  # down
        draw.polygon([(x, y), (x - 5, y - 8), (x + 5, y - 8)], fill="#333333")
    elif end[0] > start[0]:  # right
        draw.polygon([(x, y), (x - 8, y - 5), (x - 8, y + 5)], fill="#333333")
    elif end[0] < start[0]:  # left
        draw.polygon([(x, y), (x + 8, y - 5), (x + 8, y + 5)], fill="#333333")


def save_week2_architecture(path: Path) -> None:
    w, h = 1400, 980
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title_f, box_f, small_f = _font(22, True), _font(16, True), _font(14)

    d.text((40, 24), "ProcurePrep UG — Week 2 SME Reorder Architecture", fill="#111", font=title_f)

    _draw_box(d, (520, 80, 880, 140), "SME Procurement Officer", box_f, fill="#E8F1FB")
    _arrow(d, (700, 140), (700, 170))
    _draw_box(d, (500, 170, 900, 230), "Application Layer  (main.py)", box_f, fill="#F4F4F4")
    _arrow(d, (700, 230), (700, 260))
    _draw_box(d, (430, 260, 970, 340), "Deterministic Reorder Gate   current_stock < reorder_threshold", box_f, fill="#FFF4D6")

    _arrow(d, (520, 340), (280, 400))
    _arrow(d, (880, 340), (1120, 400))
    _draw_box(d, (80, 400, 480, 470), "Stock sufficient  —  skip LLM", box_f, fill="#E6F4EA")
    _draw_box(d, (920, 400, 1320, 470), "Deficit  —  load quotations.csv", box_f, fill="#FDECEA")

    _arrow(d, (1120, 470), (1120, 510))
    _draw_box(d, (880, 510, 1320, 610), "0 quotes: deterministic alert\n≥1 quote: Gemini 3.5 Flash-Lite", small_f, fill="#FDECEA")

    _arrow(d, (1120, 610), (1120, 650))
    _draw_box(d, (840, 650, 1360, 730), "Objective quotation comparison  (Prompt v1.1)", box_f, fill="#E8F1FB")

    _arrow(d, (280, 470), (280, 800))
    _arrow(d, (1120, 730), (700, 800))
    _arrow(d, (280, 800), (700, 800))
    _draw_box(d, (430, 800, 970, 900), "Human store manager review  —  no autonomous award", box_f, fill="#EDE7F6")

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "JPEG", quality=92)
    print(f"  Wrote {path.relative_to(PROJECT_ROOT)}")


def save_week3_architecture(path: Path) -> None:
    w, h = 1500, 1040
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title_f, box_f, small_f = _font(22, True), _font(16, True), _font(14)

    d.text((40, 24), "ProcurePrep UG — Week 3 RAG Architecture", fill="#111", font=title_f)

    _draw_box(d, (40, 90, 420, 160), "knowledge/*.txt  (12 docs)", box_f, fill="#E8F1FB")
    _arrow(d, (430, 125), (500, 125))
    _draw_box(d, (500, 90, 780, 160), "ingest.py  (chunk)", box_f, fill="#F4F4F4")
    _arrow(d, (780, 125), (850, 125))
    _draw_box(d, (850, 90, 1160, 160), "embed_store.py", box_f, fill="#F4F4F4")
    _arrow(d, (1160, 125), (1230, 125))
    _draw_box(d, (1230, 80, 1460, 200), "ChromaDB\nprocureprep_corpus", small_f, fill="#FFF4D6")

    _draw_box(d, (40, 260, 420, 330), "inventory.csv + quotations.csv", box_f, fill="#E8F1FB")
    _arrow(d, (230, 330), (230, 370))
    _draw_box(d, (40, 370, 420, 460), "Deterministic reorder gate", box_f, fill="#FFF4D6")
    _arrow(d, (230, 460), (230, 500))
    _draw_box(d, (40, 500, 420, 590), "Retrieval query  (item + quotes)", box_f, fill="#FDECEA")

    _arrow(d, (420, 545), (1230, 200))
    _arrow(d, (1345, 200), (1345, 500))
    _draw_box(d, (1120, 500, 1460, 590), "Top-k chunks + scores", box_f, fill="#FFF4D6")

    _arrow(d, (230, 590), (230, 660))
    _arrow(d, (1290, 590), (400, 690))
    _draw_box(d, (40, 660, 720, 750), "RELEVANT KNOWLEDGE + quotation payload", box_f, fill="#E6F4EA")
    _arrow(d, (380, 750), (380, 800))
    _draw_box(d, (40, 800, 720, 880), "Gemini + prompt_rag_v1.0  (SOURCES USED)", box_f, fill="#E8F1FB")
    _arrow(d, (380, 880), (380, 930))
    _draw_box(d, (40, 930, 720, 1000), "Human store manager review", box_f, fill="#EDE7F6")

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "JPEG", quality=92)
    print(f"  Wrote {path.relative_to(PROJECT_ROOT)}")


def save_week4_architecture(path: Path) -> None:
    w, h = 1560, 1080
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title_f, box_f, small_f = _font(22, True), _font(15, True), _font(13)

    d.text((40, 20), "ProcurePrep UG — Week 4 Tool-Calling Architecture", fill="#111", font=title_f)

    _draw_box(d, (560, 70, 1000, 130), "SME user task (natural language)", box_f, fill="#E8F1FB")
    _arrow(d, (780, 130), (780, 165))
    _draw_box(d, (480, 165, 1080, 245), "Gemini 3.5 Flash-Lite   agent_orchestrator.py   AFC disabled", small_f, fill="#F4F4F4")

    _arrow(d, (780, 245), (780, 285))
    _draw_box(d, (520, 285, 1040, 345), "Approved tool allow-list  (dispatch_tool)", box_f, fill="#FFF4D6")

    # four tools
    y = 400
    tools = [
        (40, "get_reorder_status", "Read-only", "#E6F4EA"),
        (420, "get_supplier_quotations", "Read current CSV data", "#E6F4EA"),
        (800, "create_requisition_draft", "Low-risk DRAFT record", "#E8F1FB"),
        (1180, "route_requisition_for_approval", "HIGH IMPACT — gated", "#FDECEA"),
    ]
    for x, name, note, color in tools:
        _arrow(d, (780, 345), (x + 150, y))
        _draw_box(d, (x, y, x + 340, y + 90), f"{name}\n{note}", small_f, fill=color)

    _arrow(d, (210, 490), (210, 540))
    _arrow(d, (590, 490), (590, 540))
    _draw_box(d, (40, 540, 760, 610), "inventory.csv   /   quotations.csv", box_f, fill="#F4F4F4")

    _arrow(d, (970, 490), (970, 540))
    _draw_box(d, (800, 540, 1140, 610), "requisitions.json  DRAFT", box_f, fill="#E8F1FB")

    _arrow(d, (1350, 490), (1350, 540))
    _draw_box(d, (1180, 540, 1520, 640), "CLI human y/n confirm", box_f, fill="#EDE7F6")
    _arrow(d, (1350, 640), (1350, 680))
    _draw_box(d, (1180, 680, 1520, 780), "Decline: HUMAN_DECLINED\nAccept: PENDING_APPROVAL", small_f, fill="#FDECEA")

    _draw_box(d, (40, 680, 1140, 800),
              "Deterministic validation (not LLM): missing params, unknown tool, expired quote,\n"
              "wrong-item quotation, malformed record, invalid status, unavailable Gemini API",
              small_f, fill="#FFF4D6")

    _arrow(d, (590, 800), (590, 860))
    _draw_box(d, (280, 860, 900, 960), "Human store manager  —  no autonomous purchase or award", box_f, fill="#EDE7F6")

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "JPEG", quality=92)
    print(f"  Wrote {path.relative_to(PROJECT_ROOT)}")


# ---------------------------------------------------------------------------
# Week 1-style progress report Word document
# ---------------------------------------------------------------------------

NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def _run(p, text, *, size=11, bold=False, italic=False, color=None, name="Calibri"):
    run = p.add_run(text)
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rPr.insert(0, rFonts)
    return run


def build_week1_style_progress(title_week: str, md_path: Path, out_path: Path) -> None:
    """Build a short progress report that mirrors WEEK1 structure."""
    from build_submission_docx import add_rich_paragraph, _parse_inline

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, "BSE4104 AI-Native & Agentic Engineering Capstone", size=16, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, title_week, size=14, bold=True)

    md = md_path.read_text(encoding="utf-8")
    # Skip the markdown H1/H2 title block; keep from Project: onward
    body = md
    for marker in ("**Project:**", "### 1."):
        if marker in md:
            body = md[md.find(marker):]
            break
    parse_markdown_to_doc(doc, body, skip_first_h1=False)
    add_page_numbers(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"  Created {out_path.relative_to(PROJECT_ROOT)}")


# ---------------------------------------------------------------------------
# Folder layout
# ---------------------------------------------------------------------------

WEEK2_MOVES = [
    ("Model_Selection_Note.docx", "ProcurePrep UG Model Selection Note.docx"),
    ("Prompt_Specification_v1.0.docx", "ProcurePrep UG Prompt Specification v1.0.docx"),
    ("Prompt_Specification_v1.1_and_Version_History.docx", "ProcurePrep UG Prompt Specification v1.1 and Version History.docx"),
    ("Week2_10Case_Evaluation_Table.docx", "ProcurePrep UG 10-Case Evaluation Table.docx"),
]

WEEK3_MOVES = [
    ("Corpus_Register.docx", "ProcurePrep UG Corpus Register.docx"),
    ("RAG_Architecture.docx", "ProcurePrep UG RAG Architecture.docx"),
    ("Week3_RAG_Evaluation.docx", "ProcurePrep UG RAG Evaluation.docx"),
    ("Sample_Grounded_Comparison.docx", "ProcurePrep UG Sample Grounded Comparison.docx"),
]


def _relocate(src_dir: Path, dest_dir: Path, mapping: list[tuple[str, str]]) -> None:
    dest_dir.mkdir(parents=True, exist_ok=True)
    for old, new in mapping:
        src = src_dir / old
        dest = dest_dir / new
        if src.exists():
            if dest.exists() and dest.resolve() != src.resolve():
                dest.unlink()
            shutil.move(str(src), str(dest))
            print(f"  Moved {old} -> {dest.relative_to(PROJECT_ROOT)}")
        elif dest.exists():
            print(f"  Already in place: {dest.relative_to(PROJECT_ROOT)}")
        else:
            print(f"  MISSING: {src}")


def arrange_week_folders() -> None:
    _relocate(SUBMISSION / "week2", SUBMISSION / "WEEK2", WEEK2_MOVES)
    _relocate(SUBMISSION / "week3", SUBMISSION / "WEEK3", WEEK3_MOVES)

    for leftover_name in ("Week2_Progress_Report.docx", "Week3_Progress_Report.docx"):
        for folder in (SUBMISSION / "week2", SUBMISSION / "week3", SUBMISSION / "WEEK2", SUBMISSION / "WEEK3"):
            leftover = folder / leftover_name
            if leftover.exists():
                leftover.unlink()
                print(f"  Removed old long report {leftover.relative_to(PROJECT_ROOT)}")

    for old in (SUBMISSION / "week2", SUBMISSION / "week3"):
        if old.exists() and not any(old.iterdir()):
            old.rmdir()
            print(f"  Removed empty {old.relative_to(PROJECT_ROOT)}")


def build_week4_docs() -> None:
    week4 = SUBMISSION / "WEEK4"
    week4.mkdir(parents=True, exist_ok=True)

    # Tool catalogue
    doc = Document()
    add_title_block(doc, "Tool Catalogue")
    parse_markdown_to_doc(doc, (PROJECT_ROOT / "docs/tool-catalogue.md").read_text(encoding="utf-8"))
    add_page_numbers(doc)
    out = week4 / "ProcurePrep UG Tool Catalogue.docx"
    doc.save(str(out))
    print(f"  Created {out.relative_to(PROJECT_ROOT)}")

    # Agent architecture write-up
    doc = Document()
    add_title_block(doc, "Updated Agent Architecture")
    parse_markdown_to_doc(
        doc,
        (PROJECT_ROOT / "docs/architecture/agent-architecture.md").read_text(encoding="utf-8"),
        render_diagrams=False,
    )
    add_page_numbers(doc)
    out = week4 / "ProcurePrep UG Agent Architecture.docx"
    doc.save(str(out))
    print(f"  Created {out.relative_to(PROJECT_ROOT)}")

    # Demo
    demo_md = PROJECT_ROOT / "results/week4_tool_demo.md"
    if demo_md.exists():
        doc = Document()
        add_title_block(doc, "Tool Demonstration")
        parse_markdown_to_doc(doc, demo_md.read_text(encoding="utf-8"))
        add_page_numbers(doc)
        out = week4 / "ProcurePrep UG Tool Demonstration.docx"
        doc.save(str(out))
        print(f"  Created {out.relative_to(PROJECT_ROOT)}")

    # Tests
    tests_md = PROJECT_ROOT / "results/week4_tool_tests.md"
    if tests_md.exists():
        doc = Document()
        add_title_block(doc, "Failure and Authorization Tests")
        parse_markdown_to_doc(doc, tests_md.read_text(encoding="utf-8"))
        add_page_numbers(doc)
        out = week4 / "ProcurePrep UG Failure Authorization Tests.docx"
        doc.save(str(out))
        print(f"  Created {out.relative_to(PROJECT_ROOT)}")


def main() -> None:
    print("Arranging WEEK2 / WEEK3 folders to match WEEK1...")
    arrange_week_folders()

    print("\nGenerating architecture diagrams...")
    save_week2_architecture(SUBMISSION / "WEEK2" / "SME_Reorder_Architecture.jpg")
    save_week3_architecture(SUBMISSION / "WEEK3" / "RAG_Architecture.jpg")
    save_week4_architecture(SUBMISSION / "WEEK4" / "Agent_Architecture.jpg")

    print("\nBuilding Week 1-style progress reports...")
    build_week1_style_progress(
        "Week 2 Progress Report",
        PROJECT_ROOT / "docs/weekly-reports/week2-progress.md",
        SUBMISSION / "WEEK2" / "BSE4104 Week 2 Progress Report.docx",
    )
    build_week1_style_progress(
        "Week 3 Progress Report",
        PROJECT_ROOT / "docs/weekly-reports/week3-progress.md",
        SUBMISSION / "WEEK3" / "BSE4104 Week 3 Progress Report.docx",
    )
    build_week1_style_progress(
        "Week 4 Progress Report",
        PROJECT_ROOT / "docs/weekly-reports/week4.md",
        SUBMISSION / "WEEK4" / "BSE4104 Week 4 Progress Report.docx",
    )

    print("\nBuilding Week 4 assignment documents...")
    build_week4_docs()
    print("\nDone.")


if __name__ == "__main__":
    main()
