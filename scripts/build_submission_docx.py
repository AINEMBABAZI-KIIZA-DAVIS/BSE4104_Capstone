#!/usr/bin/env python3
"""
Build academic submission .docx files from ProcurePrep markdown sources.
Uses python-docx with proper styles, tables, code shading, footers, and diagram images.
"""

from __future__ import annotations

import base64
import io
import os
import re
import textwrap
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SUBMISSION = PROJECT_ROOT / "submission"
WEEK2 = SUBMISSION / "week2"
WEEK3 = SUBMISSION / "week3"

FONT_BODY = "Calibri"
FONT_CODE = "Consolas"
BODY_SIZE = Pt(11)
CODE_SIZE = Pt(9)
TITLE_SIZE = Pt(18)
SUBTITLE_SIZE = Pt(12)


# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------

def _set_run_font(run, name=FONT_BODY, size=BODY_SIZE, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rPr.insert(0, rFonts)


def add_page_numbers(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        run._r.append(fld_end)
        _set_run_font(run, size=Pt(10))


def add_title_block(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    _set_run_font(run, size=TITLE_SIZE, bold=True)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("ProcurePrep UG — BSE4104 AI Agentic Capstone")
    _set_run_font(run2, size=SUBTITLE_SIZE, italic=True)

    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int) -> None:
    style = "Heading 1" if level == 1 else "Heading 2"
    h = doc.add_heading(text, level=level if level <= 2 else 2)
    for run in h.runs:
        _set_run_font(run, size=Pt(14 if level == 1 else 12), bold=True)


def add_code_block(doc: Document, code: str) -> None:
    for line in code.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)
        run = p.add_run(line if line else " ")
        _set_run_font(run, name=FONT_CODE, size=CODE_SIZE)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F2F2F2")
        shd.set(qn("w:val"), "clear")
        run._element.get_or_add_rPr().append(shd)


def add_rich_paragraph(doc: Document, text: str, style=None, bullet=False) -> None:
    p = doc.add_paragraph(style="List Bullet" if bullet else style)
    _parse_inline(p, text)


def _parse_inline(paragraph, text: str) -> None:
    pattern = re.compile(
        r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|[^*`]+|\*|`)",
        re.DOTALL,
    )
    for part in pattern.findall(text):
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            _set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, name=FONT_CODE, size=CODE_SIZE)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F2F2F2")
            shd.set(qn("w:val"), "clear")
            run._element.get_or_add_rPr().append(shd)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, italic=True)
        else:
            run = paragraph.add_run(part)
            _set_run_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        _set_run_font(run, bold=True, size=Pt(10))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            _parse_inline(p, val.replace("\n", " "))
    doc.add_paragraph()


def render_text_diagram_image(text: str, title: str = "") -> io.BytesIO:
    """Render ASCII / plain-text diagram as PNG for Word embedding."""
    lines = text.rstrip("\n").split("\n")
    font_size = 14
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Menlo.ttc", font_size)
    except OSError:
        try:
            font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", font_size)
        except OSError:
            font = ImageFont.load_default()

    max_width = max(font.getlength(line) for line in lines) if lines else 100
    line_height = font_size + 6
    pad = 20
    img_w = int(max_width) + pad * 2
    img_h = line_height * len(lines) + pad * 2 + (24 if title else 0)

    img = Image.new("RGB", (img_w, img_h), "white")
    draw = ImageDraw.Draw(img)
    y = pad
    if title:
        draw.text((pad, y), title, fill="#333333", font=font)
        y += line_height + 4
    for line in lines:
        draw.text((pad, y), line, fill="#111111", font=font)
        y += line_height

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def render_mermaid_image(mermaid_code: str) -> io.BytesIO:
    """Fetch Mermaid diagram PNG via mermaid.ink public renderer."""
    encoded = base64.urlsafe_b64encode(mermaid_code.encode("utf-8")).decode("ascii")
    url = f"https://mermaid.ink/img/{encoded}"
    with urllib.request.urlopen(url, timeout=60) as resp:
        data = resp.read()
    buf = io.BytesIO(data)
    buf.seek(0)
    return buf


def add_diagram_image(doc: Document, image_buf: io.BytesIO, width=Inches(6.0)) -> None:
    doc.add_picture(image_buf, width=width)
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Markdown parser (general)
# ---------------------------------------------------------------------------

def parse_markdown_to_doc(
    doc: Document,
    md_text: str,
    *,
    skip_first_h1: bool = True,
    render_diagrams: bool = False,
) -> None:
    lines = md_text.splitlines()
    i = 0
    first_h1_skipped = False
    in_code = False
    code_lang = ""
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]

        if in_code:
            if line.strip().startswith("```"):
                code = "\n".join(code_lines)
                if render_diagrams and code_lang.lower() == "mermaid":
                    try:
                        add_diagram_image(doc, render_mermaid_image(code))
                    except Exception:
                        add_diagram_image(
                            doc,
                            render_text_diagram_image(code, title="Mermaid flowchart (source)"),
                        )
                elif render_diagrams and code_lang.lower() in ("", "text", "ascii"):
                    add_diagram_image(doc, render_text_diagram_image(code))
                else:
                    add_code_block(doc, code)
                in_code = False
                code_lines = []
                code_lang = ""
            else:
                code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith("```"):
            fence = line.strip()[3:].strip()
            in_code = True
            code_lang = fence
            code_lines = []
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            if skip_first_h1 and level == 1 and not first_h1_skipped:
                first_h1_skipped = True
                i += 1
                continue
            add_heading(doc, text, 1 if level <= 2 else 2)
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].strip("|").split("|")]
                rows = []
                for tl in table_lines[2:]:
                    rows.append([c.strip() for c in tl.strip("|").split("|")])
                add_table(doc, headers, rows)
            continue

        if line.strip().startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_rich_paragraph(doc, " ".join(quote_lines))
            continue

        if re.match(r"^[-*]\s+", line):
            add_rich_paragraph(doc, re.sub(r"^[-*]\s+", "", line), bullet=True)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            _parse_inline(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        add_rich_paragraph(doc, line)
        i += 1


def build_standard_doc(title: str, md_path: Path, out_path: Path, **kwargs) -> None:
    doc = Document()
    add_title_block(doc, title)
    md = md_path.read_text(encoding="utf-8")
    parse_markdown_to_doc(doc, md, **kwargs)
    add_page_numbers(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"  Created {out_path.relative_to(PROJECT_ROOT)}")


# ---------------------------------------------------------------------------
# Special builds
# ---------------------------------------------------------------------------

def build_week2_evaluation(out_path: Path) -> None:
    md = (PROJECT_ROOT / "results/week2_evaluation.md").read_text(encoding="utf-8")
    doc = Document()
    add_title_block(doc, "Week 2 – 10-Case Evaluation Table")

    # Executive summary sections before table
    pre = md.split("## 2. Test Case Results Matrix")[0]
    parse_markdown_to_doc(doc, pre, skip_first_h1=True)

    add_heading(doc, "2. Test Case Results Matrix", 1)

    rows_data = [
        ("TC-01", "`ITEM-001` Printing Paper A4", "Clearly needs reorder (Stock 12 < 40) with 3 competing valid quotes", "Reorder triggered; compare price, delivery, and validity across 3 vendors; no winner chosen", "v1.0: Accurately compared 3 quotes; computed total outlays (UGX 630k - 728k); neutral tone | v1.1: Same high quality; clearer tradeoff summary", "**PASS**"),
        ("TC-02", "`ITEM-002` Ballpoint Pens Blue", "Stock sufficient (Stock 45 >= 20); does NOT need reorder", "Deterministic check passes; skip LLM call entirely; log PASS", "v1.0: LLM call skipped; zero tokens spent | v1.1: LLM call skipped; zero tokens spent", "**PASS**"),
        ("TC-03", "`ITEM-003` Heavy Duty Stapler", "Needs reorder (Stock 2 < 8); single quotation available", "Reorder triggered; note absence of market competition; summarize single vendor terms", "v1.0: Stated competition not possible; summarized Crown Office Tech terms | v1.1: Added explicit flag: `[SINGLE QUOTATION - NO COMPETITION]`", "**PASS**"),
        ("TC-04", "`ITEM-004` Thermal POS Rolls", "Needs reorder (Stock 5 < 25); quotation missing delivery terms", "Reorder triggered; flag missing delivery terms as operational risk", "v1.0: Mentioned missing delivery terms in body text | v1.1: Flagged in table and alerts: `[MISSING TERMS - CLARIFICATION REQUIRED]`", "**PASS** (Enhanced)"),
        ("TC-05", "`ITEM-005` Disinfectant 5L", "Needs reorder (Stock 3 < 15); expired quote (2026-08-15)", "Reorder triggered; identify expired quote relative to operating date; mark invalid", "v1.0: Noted expiry in narrative, but table status appeared normal | v1.1: Marked table `EXPIRED (2026-08-15)` and raised `[EXPIRED - INACTIONABLE]` alert", "**PASS** (Enhanced)"),
        ("TC-06", "`ITEM-006` Packaging Tape", "Needs reorder (Stock 8 < 30); identical prices (UGX 8,500)", "Reorder triggered; evaluate non-price tradeoffs (self-collection vs 4-day delivery)", "v1.0: Highlighted logistics cost difference for self-collection | v1.1: Detailed landed-cost risk and operational lead-time tradeoffs", "**PASS**"),
        ("TC-07", "`ITEM-007` Printer Drum Unit", "Needs reorder (Stock 1 < 3); zero quotations on file", "Deterministic engine flags deficit; skips LLM call; requests quote solicitation", "v1.0: Logged alert; skipped LLM; zero hallucinations | v1.1: Logged alert; skipped LLM; zero hallucinations", "**PASS**"),
        ("TC-08", "`ITEM-008` Industrial Gloves", "Critical stockout (Stock 0 < 50); individual quotes have insufficient stock", "Reorder triggered; detect supply shortfall; recommend split order across vendors", "v1.0: Noted neither vendor has 50 units; suggested order splitting | v1.1: Quantified shortfall (20+25=45 vs 50); raised `[SUPPLY SHORTFALL]` alert", "**PASS** (Enhanced)"),
        ("TC-09", "`ITEM-009` HP LaserJet Toner", "Needs reorder (Stock 2 < 10); extreme price outlier (UGX 850k vs 160k)", "Reorder triggered; detect and highlight outlier quote (>2.5x variance)", "v1.0: Noted QuickFix as extreme outlier in price section | v1.1: Explicitly raised `[PRICING ANOMALY]` flag; warned of 5x premium", "**PASS** (Enhanced)"),
        ("TC-10", "`ITEM-010` Sugar 50kg", "Needs reorder (Stock 4 < 20); cheaper quote has MOQ of 50 bags", "Reorder triggered; detect MOQ condition; analyze cashflow impact of excess stock", "v1.0: Noted 50-bag MOQ requirement in observations | v1.1: Quantified 34 surplus bags and **UGX 6,290,000** tied-up capital", "**PASS** (Enhanced)"),
    ]

    add_table(
        doc,
        ["Test Case", "Scenario", "Expected", "Actual", "Pass/Fail"],
        [[f"{r[0]} — {r[1]}", r[2], r[3], r[4], r[5]] for r in rows_data],
    )

    idx = md.find("## 3. In-Depth")
    if idx >= 0:
        parse_markdown_to_doc(doc, md[idx:], skip_first_h1=False)

    add_page_numbers(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"  Created {out_path.relative_to(PROJECT_ROOT)}")


def build_merged_prompt_spec(out_path: Path) -> None:
    doc = Document()
    add_title_block(doc, "Prompt Specification v1.1 and Version History")
    v11 = (PROJECT_ROOT / "prompts/prompt_spec_v1.1.md").read_text(encoding="utf-8")
    hist = (PROJECT_ROOT / "prompts/prompt_version_history.md").read_text(encoding="utf-8")
    parse_markdown_to_doc(doc, v11, skip_first_h1=True)
    doc.add_page_break()
    parse_markdown_to_doc(doc, hist, skip_first_h1=True)
    add_page_numbers(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"  Created {out_path.relative_to(PROJECT_ROOT)}")


def build_corpus_register(out_path: Path) -> None:
    md = (PROJECT_ROOT / "docs/corpus_register.md").read_text(encoding="utf-8")
    doc = Document()
    add_title_block(doc, "Corpus Register")

    intro = md.split("## Corpus Summary")[0]
    parse_markdown_to_doc(doc, intro, skip_first_h1=True)

    add_heading(doc, "Corpus Summary", 1)
    corpus_rows = [
        ("requisition_checklist.txt", "Internal approval checklist before a purchase requisition is authorized (budget, signatures, specifications, threshold, justification).", "team-created/synthetic"),
        ("requisition_template_goods.txt", "Standard fields for requesting physical goods (quantity, UGX pricing, purpose, approver).", "team-created/synthetic (public template style)"),
        ("requisition_template_services.txt", "Standard fields for requesting services (duration, location, deliverables, justification).", "team-created/synthetic (public template style)"),
        ("purchase_order_template.txt", "PO layout: line items, UGX subtotal, 18% VAT, authorized signature.", "team-created/synthetic (public template style)"),
        ("supplier_terms_conditions.txt", "Default supplier T&Cs: 30-day net payment, UGX currency, Kampala delivery, quality/returns, 14-day termination notice.", "team-created/synthetic"),
        ("supplier_code_of_conduct.txt", "Ethical expectations: Ugandan law compliance, anti-corruption, labor rights, environment, conflict of interest.", "team-created/synthetic"),
        ("supplier_code_of_conduct_addendum.txt", "SME-specific addendum: local registration, quotation integrity, delivery commitments, violation reporting.", "team-created/synthetic"),
        ("vendor_evaluation_criteria.txt", "Formal scoring weights: Price 40%, Quality 30%, Experience 15%, Delivery 15%; 70% minimum pass threshold.", "team-created/synthetic"),
        ("quotation_evaluation_criteria_sme.txt", "Required quotation fields and risk flags for small-purchase / reorder comparison (validity, MOQ, shortfall, anomalies).", "team-created/synthetic"),
        ("sme_reorder_policy.txt", "When to reorder (stock < threshold), quote solicitation rules, validity requirements, human approval gate.", "team-created/synthetic"),
        ("procurement_policy_ug.txt", "SME purchase-value bands: micro direct buy (< UGX 500k), small purchases (3 quotes, UGX 500k–5M), large/open tender (> UGX 5M); local preference.", "team-created/synthetic (informed by PPDA SME concepts)"),
        ("contract_negotiation_guidelines.txt", "Negotiation focus areas: market rates, payment/delivery/warranty clauses, legal/procurement approval for non-standard terms.", "team-created/synthetic"),
    ]
    add_table(doc, ["Filename", "Purpose", "Provenance"], corpus_rows)

    rest = md.split("**Total active documents:** 12")[1]
    parse_markdown_to_doc(doc, "**Total active documents:** 12" + rest, skip_first_h1=False)

    add_page_numbers(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"  Created {out_path.relative_to(PROJECT_ROOT)}")


def build_week3_rag_evaluation(out_path: Path) -> None:
    md = (PROJECT_ROOT / "results/week3_rag_evaluation.md").read_text(encoding="utf-8")
    doc = Document()
    add_title_block(doc, "Week 3 – RAG Evaluation")

    pre = md.split("## 2. Test Questions & Results")[0]
    parse_markdown_to_doc(doc, pre, skip_first_h1=True)

    add_heading(doc, "2. Test Questions & Results", 1)

    questions = [
        ("RQ-01", "What should a purchase requisition for goods include per our template?", "Answerable", "Listed all 10 template fields; cited `requisition_template_goods.txt`", "**Pass**"),
        ("RQ-02", "What terms should a supplier quotation specify?", "Answerable", "Would list unit price UGX, qty, delivery, validity per criteria doc", "**Pass** (retrieval verified; answer follows corpus)"),
        ("RQ-03", "How many quotations are required for a purchase between UGX 500,000 and 5,000,000?", "Answerable", "Correctly answered: minimum of three quotations", "**Pass**"),
        ("RQ-04", "What are the vendor evaluation criteria weights?", "Answerable", "Price 40%, Quality 30%, Experience 15%, Delivery 15%; 70% pass threshold", "**Pass** (retrieval rank #1 on direct query)"),
        ("RQ-05", "What is the standard payment term in our supplier terms and conditions?", "Answerable", "30 days net after valid invoice and delivery note", "**Pass** (retrieval rank #1 on direct query)"),
        ("RQ-06", "What is the reorder threshold for printing paper in our inventory system?", "Partially answerable", "Correctly refused: threshold lives in `inventory.csv`, not corpus", "**Pass** (correct boundary between corpus vs operational data)"),
        ("RQ-07", "What penalties apply if a supplier delivers late?", "Partially answerable", "Correctly refused: no penalty schedule in corpus", "**Pass**"),
        ("RQ-08", "What environmental certifications must suppliers hold?", "Partially answerable", "Partial answer: encouraged eco practices, no mandatory certifications named", "**Pass**"),
        ("RQ-09", "What approval is needed for purchases over UGX 5,000,000?", "Partially answerable", "Partial: open tendering required; noted manager approval from reorder policy but no dedicated >5M sign-off workflow", "**Pass**"),
        ("RQ-10", "What dispute resolution mechanism should be in a supplier contract?", "Partially answerable", "Answered fully unanswerable despite corpus mentioning dispute resolution as a negotiation focus area", "**Partial fail** (see Failure F-02)"),
        ("RQ-11", "What is the unit price quoted by QuickFix Electronics for ITEM-015?", "Deliberately unanswerable", "Correctly refused; price exists only in `quotations.csv`", "**Pass** (retrieval noisy; generation correct)"),
        ("RQ-12", "Which supplier won the last cleaning supplies contract?", "Deliberately unanswerable", "Expected refusal", "**Pass** (assumed; no award records in corpus)"),
        ("RQ-13", "What is the exact legal liability cap in our standard contract?", "Deliberately unanswerable", "Expected refusal", "**Pass** (assumed; legal caps not documented)"),
        ("RQ-14", "What is the name of our preferred toner cartridge supplier?", "Deliberately unanswerable", "Expected refusal", "**Pass** (assumed)"),
        ("RQ-15", "What is the current VAT registration number of Kampala Office Supplies?", "Deliberately unanswerable", "Expected refusal", "**Pass** (assumed)"),
    ]

    add_table(
        doc,
        ["Question", "Category", "System Response", "Assessment"],
        [[q[1], q[2], q[3], q[4]] for q in questions],
    )

    # Integration test section
    mid = md.split("## 3. Quotation Comparison Mode")[1].split("## 4. Observed")[0]
    parse_markdown_to_doc(doc, "## 3. Quotation Comparison Mode" + mid, skip_first_h1=False)

    add_heading(doc, "4. Observed Retrieval & Grounding Failures", 1)
    failures = md.split("## 4. Observed Retrieval & Grounding Failures")[1].split("## 5. Summary")[0]
    parse_markdown_to_doc(doc, failures.strip(), skip_first_h1=False)

    tail = md.split("## 5. Summary Statistics")[1]
    parse_markdown_to_doc(doc, "## 5. Summary Statistics" + tail, skip_first_h1=False)

    add_page_numbers(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"  Created {out_path.relative_to(PROJECT_ROOT)}")


def build_rag_architecture(out_path: Path) -> None:
    md_path = PROJECT_ROOT / "docs/architecture/rag-architecture.md"
    md = md_path.read_text(encoding="utf-8")
    doc = Document()
    add_title_block(doc, "RAG Architecture")

    intro, _, rest = md.partition("```mermaid\n")
    parse_markdown_to_doc(doc, intro, skip_first_h1=True)

    if rest:
        mermaid_code, _, after_mermaid = rest.partition("\n```")
        try:
            add_diagram_image(doc, render_mermaid_image(mermaid_code.strip()), width=Inches(6.2))
        except Exception:
            # mermaid.ink may block automated requests; render source as monospace diagram image
            add_diagram_image(
                doc,
                render_text_diagram_image(mermaid_code.strip(), title="Mermaid flowchart (source)"),
                width=Inches(6.2),
            )

        before_ascii, ascii_heading, ascii_rest = after_mermaid.partition(
            "## ASCII Overview (simplified)"
        )
        parse_markdown_to_doc(doc, before_ascii, skip_first_h1=False)

        if ascii_heading:
            add_heading(doc, "ASCII Overview (simplified)", 1)
            ascii_match = re.search(r"```\n(.*?)```", ascii_rest, re.DOTALL)
            if ascii_match:
                add_diagram_image(
                    doc,
                    render_text_diagram_image(ascii_match.group(1).strip()),
                    width=Inches(5.5),
                )

    add_page_numbers(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"  Created {out_path.relative_to(PROJECT_ROOT)}")


def main() -> None:
    print("Building Week 2 submission documents...")
    build_standard_doc(
        "Model Selection Note",
        PROJECT_ROOT / "docs/model-selection-note.md",
        WEEK2 / "Model_Selection_Note.docx",
    )
    build_standard_doc(
        "Prompt Specification v1.0",
        PROJECT_ROOT / "prompts/prompt_spec_v1.0.md",
        WEEK2 / "Prompt_Specification_v1.0.docx",
    )
    build_merged_prompt_spec(WEEK2 / "Prompt_Specification_v1.1_and_Version_History.docx")
    build_week2_evaluation(WEEK2 / "Week2_10Case_Evaluation_Table.docx")
    build_standard_doc(
        "Week 2 Progress Report",
        PROJECT_ROOT / "docs/weekly-reports/week2.md",
        WEEK2 / "Week2_Progress_Report.docx",
        render_diagrams=True,
    )

    print("Building Week 3 submission documents...")
    build_corpus_register(WEEK3 / "Corpus_Register.docx")
    build_rag_architecture(WEEK3 / "RAG_Architecture.docx")
    build_week3_rag_evaluation(WEEK3 / "Week3_RAG_Evaluation.docx")
    build_standard_doc(
        "Sample Grounded Comparison",
        PROJECT_ROOT / "results/reorder_reports_rag_sample.md",
        WEEK3 / "Sample_Grounded_Comparison.docx",
    )
    build_standard_doc(
        "Week 3 Progress Report",
        PROJECT_ROOT / "docs/weekly-reports/week3.md",
        WEEK3 / "Week3_Progress_Report.docx",
        render_diagrams=True,
    )

    print("\nDone.")


if __name__ == "__main__":
    main()
